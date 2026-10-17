import io

import docx
import fitz
import pytest

from app.core.errors import ValidationFailed
from app.services import text_processing as tp


def make_pdf(pages: list[str]) -> bytes:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(50, 50, 550, 800), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Leave Policy", level=1)
    d.add_paragraph("Employees receive 25 days of annual leave.")
    d.add_paragraph("Carry-over", style="List Bullet")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Grade"
    table.rows[0].cells[1].text = "Days"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_detect_type_validates_extension_and_magic_bytes():
    assert tp.detect_type("a.pdf", b"%PDF-1.7 ...") == "pdf"
    assert tp.detect_type("notes.MD", b"# hi") == "md"
    assert tp.detect_type("a.txt", b"hello") == "txt"
    assert tp.detect_type("a.docx", make_docx()) == "docx"
    with pytest.raises(ValidationFailed):
        tp.detect_type("a.exe", b"MZ")
    with pytest.raises(ValidationFailed):
        tp.detect_type("fake.pdf", b"not a pdf")
    with pytest.raises(ValidationFailed):
        tp.detect_type("fake.docx", b"PK\x03\x04garbage")
    with pytest.raises(ValidationFailed):
        tp.detect_type("bin.txt", b"abc\x00def")


def test_sanitize_filename_strips_paths_and_control_chars():
    assert tp.sanitize_filename("../../etc/passwd") == "passwd"
    assert tp.sanitize_filename("C:\\Users\\x\\report\x00.pdf") == "report.pdf"
    assert tp.sanitize_filename("") == "untitled"


def test_extract_pdf_keeps_page_numbers():
    pages = tp.extract("pdf", make_pdf(["First page about apples.", "Second page about pears."]), max_pages=10)
    assert [p.page for p in pages] == [1, 2]
    assert "apples" in pages[0].text and "pears" in pages[1].text


def test_extract_pdf_page_limit():
    with pytest.raises(ValidationFailed, match="limit"):
        tp.extract("pdf", make_pdf(["a", "b", "c"]), max_pages=2)


def test_extract_docx_preserves_structure():
    [page] = tp.extract("docx", make_docx(), max_pages=10)
    assert page.page is None
    assert "# Leave Policy" in page.text
    assert "- Carry-over" in page.text
    assert "Grade | Days" in page.text


def test_extract_text_handles_bom_and_fallback_encoding():
    assert tp.extract("txt", "\ufeffhello".encode(), 1)[0].text == "hello"
    assert "café" in tp.extract("txt", "café".encode("cp1252"), 1)[0].text


def test_normalize_text():
    raw = "Hello\r\n\r\n\r\n\r\nwor-\nld  \t spaced\x07\n  indented"
    assert tp.normalize_text(raw) == "Hello\n\nworld spaced\nindented"


def test_split_text_respects_size_and_overlap():
    text = " ".join(f"Sentence number {i} is here." for i in range(200))
    spans = tp.split_text(text, chunk_size=300, overlap=50)
    assert len(spans) > 5
    for s, e in spans:
        assert e - s <= 300
    # consecutive chunks overlap
    for (_s1, e1), (s2, _e2) in zip(spans, spans[1:], strict=False):
        assert s2 < e1
        assert e1 - s2 <= 80
    # complete coverage
    assert spans[0][0] == 0 and spans[-1][1] == len(text)


def test_split_text_prefers_paragraph_boundaries():
    text = "A" * 150 + "\n\n" + "B" * 150 + "\n\n" + "C" * 150
    spans = tp.split_text(text, chunk_size=320, overlap=0)
    assert text[spans[0][0] : spans[0][1]].strip() == "A" * 150 + "\n\n" + "B" * 150


def test_split_text_rejects_bad_parameters():
    with pytest.raises(ValueError):
        tp.split_text("abc", chunk_size=10, overlap=10)


def test_chunk_pages_assigns_page_and_index():
    pages = [tp.PageText(1, "alpha " * 100), tp.PageText(2, "beta " * 100)]
    chunks = tp.chunk_pages(pages, chunk_size=200, overlap=20)
    assert [c.index for c in chunks] == list(range(len(chunks)))
    assert {c.page for c in chunks} == {1, 2}
    assert all("beta" not in c.text for c in chunks if c.page == 1)
