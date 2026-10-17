"""Text extraction, normalisation and chunking for uploaded documents.

Uploaded files are untrusted: types are checked by extension *and* magic bytes, sizes and
page/character counts are capped, and nothing in the content is ever executed.
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass

from app.core.errors import ValidationFailed

SUPPORTED_TYPES = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "md", ".markdown": "md"}


@dataclass
class PageText:
    page: int | None  # 1-based page number for PDFs; None when the format has no pages
    text: str


@dataclass
class TextChunk:
    index: int
    page: int | None
    text: str
    char_start: int  # offsets within the normalised page text
    char_end: int


def sanitize_filename(name: str) -> str:
    name = name.replace("\\", "/").split("/")[-1]
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    return name[:255] or "untitled"


def detect_type(filename: str, data: bytes) -> str:
    """Return the document type, validating the extension against the content."""
    lower = filename.lower()
    ext = next((e for e in SUPPORTED_TYPES if lower.endswith(e)), None)
    if ext is None:
        raise ValidationFailed(
            f"Unsupported file type for '{filename}'. Supported: PDF, DOCX, TXT, Markdown.",
            code="unsupported_file_type",
        )
    kind = SUPPORTED_TYPES[ext]
    if kind == "pdf" and not data.startswith(b"%PDF-"):
        raise ValidationFailed(f"'{filename}' does not look like a PDF file.", code="invalid_file_content")
    if kind == "docx":
        if not data.startswith(b"PK\x03\x04"):
            raise ValidationFailed(f"'{filename}' does not look like a DOCX file.", code="invalid_file_content")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if "word/document.xml" not in zf.namelist():
                    raise ValidationFailed(f"'{filename}' is not a Word document.", code="invalid_file_content")
        except zipfile.BadZipFile as exc:
            raise ValidationFailed(f"'{filename}' is a corrupt DOCX file.", code="invalid_file_content") from exc
    if kind in ("txt", "md") and b"\x00" in data[:8192]:
        raise ValidationFailed(f"'{filename}' appears to be binary, not text.", code="invalid_file_content")
    return kind


def extract_pdf(data: bytes, max_pages: int) -> list[PageText]:
    import fitz  # PyMuPDF

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise ValidationFailed("The PDF could not be opened (corrupt file?).", code="invalid_file_content") from exc
    with doc:
        if doc.needs_pass:
            raise ValidationFailed("Encrypted PDFs are not supported.", code="encrypted_pdf")
        if doc.page_count > max_pages:
            raise ValidationFailed(
                f"The PDF has {doc.page_count} pages; the limit is {max_pages}.", code="too_many_pages"
            )
        return [PageText(page=i + 1, text=page.get_text("text")) for i, page in enumerate(doc)]


def extract_docx(data: bytes) -> list[PageText]:
    import docx  # python-docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ValidationFailed("The DOCX file could not be read.", code="invalid_file_content") from exc
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append("#" * min(int(level), 6) + " " + text)
        elif style.startswith("list"):
            lines.append("- " + text)
        else:
            lines.append(text)
    for table in document.tables:
        lines.append("")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [PageText(page=None, text="\n".join(lines))]


def extract_text_file(data: bytes) -> list[PageText]:
    for encoding in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return [PageText(page=None, text=data.decode(encoding))]
        except UnicodeDecodeError:
            continue
    return [PageText(page=None, text=data.decode("cp1252", errors="replace"))]


def extract(kind: str, data: bytes, max_pages: int) -> list[PageText]:
    if kind == "pdf":
        return extract_pdf(data, max_pages)
    if kind == "docx":
        return extract_docx(data)
    return extract_text_file(data)


_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def normalize_text(text: str) -> str:
    """Normalise unicode/whitespace while keeping paragraph and list structure."""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL_RE.sub("", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # re-join words hyphenated across lines
    text = re.sub(r"[ \t ]+", " ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _snap_end(text: str, start: int, end: int) -> int:
    """Move ``end`` back to a natural boundary (paragraph, sentence, word) if one is close."""
    if end >= len(text):
        return len(text)
    window_min = start + (end - start) // 2
    for sep in ("\n\n", "\n", ". ", "? ", "! ", "; ", " "):
        pos = text.rfind(sep, window_min, end)
        if pos != -1:
            return pos + len(sep)
    return end


def _snap_start(text: str, pos: int, limit: int) -> int:
    """Move an overlap start forward to the next word boundary (not past ``limit``)."""
    if pos <= 0:
        return 0
    if text[pos - 1].isspace():
        return pos
    nxt = text.find(" ", pos, limit)
    return nxt + 1 if nxt != -1 else pos


def split_text(text: str, chunk_size: int, overlap: int) -> list[tuple[int, int]]:
    """Split ``text`` into overlapping windows of at most ``chunk_size`` characters.

    Returns ``(start, end)`` offsets. Window ends snap back to paragraph/sentence/word
    boundaries; consecutive windows overlap by roughly ``overlap`` characters.
    """
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive")
    if not 0 <= overlap < chunk_size:
        raise ValueError("overlap must be in [0, chunk_size)")
    spans: list[tuple[int, int]] = []
    start, n = 0, len(text)
    while start < n:
        end = _snap_end(text, start, min(start + chunk_size, n))
        if end <= start:
            end = min(start + chunk_size, n)
        if text[start:end].strip():
            spans.append((start, end))
        if end >= n:
            break
        nxt = _snap_start(text, max(end - overlap, start + 1), end)
        start = nxt if nxt > start else end
    return spans


def chunk_pages(pages: list[PageText], chunk_size: int, overlap: int) -> list[TextChunk]:
    """Chunk each page separately so every chunk carries a correct page number."""
    chunks: list[TextChunk] = []
    for page in pages:
        for s, e in split_text(page.text, chunk_size, overlap):
            body = page.text[s:e].strip()
            if body:
                chunks.append(TextChunk(index=len(chunks), page=page.page, text=body, char_start=s, char_end=e))
    return chunks
